#!/usr/bin/env python3
"""Collect every SwordArtsParam row and its evidence into an offline implementation dataset.

Collection coverage is separate from simulation readiness. Unknown input paths,
hit counts, formula choices, timings, and measurements remain null with blockers.
"""
import argparse
import base64
import collections
import csv
import gzip
import hashlib
import json
import re
from pathlib import Path

import openpyxl

from common import (bullet_topology, cycle_witnesses, encode, read, reachable, select_bindings,
                    verify_manifest, walk_graph)

ROOT = Path(__file__).resolve().parents[3]
TYPES = {'physical': 'Phys', 'magic': 'Mag', 'fire': 'Fire', 'lightning': 'Thun', 'holy': 'Dark'}
PARAM_FILES = ['SwordArtsParam', 'BehaviorParam_PC', 'AtkParam_Pc', 'Bullet', 'SpEffectParam',
               'FinalDamageRateParam', 'EquipParamGem', 'EquipParamWeapon', 'ReinforceParamWeapon']
EFFECT_LINKS = ['replaceSpEffectId', 'cycleOccurrenceSpEffectId', 'atkOccurrenceSpEffectId', 'spiritDeathSpEffectId']
VISUAL_PREFIXES = ('pad', 'reserved', 'vfx', 'sfx', 'seId', 'iconId', 'assetNo', 'sortId')


def projected(row):
    # Preserve all mechanics fields, including zero and sentinel values, but not cosmetic assets.
    return {k: v for k, v in row.items() if k not in ['Row ID', 'Row Name', 'ID', 'Name']
            and not k.startswith(VISUAL_PREFIXES)}


def sheet_parameters(workbook, title, source, header_row=1):
    rows = list(workbook[title].values)
    headers = rows[header_row - 1]
    out = {}
    for index, row in enumerate(rows[header_row:], header_row + 1):
        if not row or not isinstance(row[0], (int, float)) or isinstance(row[0], bool):
            continue
        key = int(row[0])
        fields = {str(k): v for k, v in zip(headers, row) if k is not None and v is not None}
        if key in out:
            if out[key]['fields'] != projected(fields) or out[key]['name'] != fields.get('Name', ''):
                raise ValueError(f'conflicting parameter ID {source}/{title}/{key}')
            out[key].setdefault('identical_duplicate_rows', []).append(index)
            continue
        out[key] = {'fields': projected(fields), 'name': fields.get('Name', ''),
                    'provenance': {'source': source, 'sheet': title, 'row': index},
                    'version_scope': 'reference_sheet_snapshot_not_verified_as_1.17_raw'}
    return out


def hks_hints(text, commit):
    lines = text.splitlines()
    lists = {}
    for name in ['IsStanceArts', 'IsRollingArts', 'IsShieldArts', 'IsHalfBlendArts']:
        start = next(i for i, line in enumerate(lines) if line.startswith('function ' + name + '('))
        end = next(i for i in range(start + 1, len(lines)) if lines[i].startswith('function '))
        body = '\n'.join(lines[start:end])
        match = re.search(r'local aow_list = \{(.*?)\}', body, re.S)
        members = []
        if match:
            for line in match.group(1).splitlines():
                member = re.match(r'\s*(\d+)\s*[,}]?', line)
                if member:
                    members.append(int(member.group(1)))
        lists[name] = {'members': members, 'line_start': start + 1, 'line_end': end}
    return {'header_declared_version': '1.15.1', 'commit_declared_version': '1.17',
            'commit_sha': commit['sha'], 'commit_message': commit['commit']['message'],
            'game_version_match': 'commit_claims_1.17; community_edition_not_pristine_game_script',
            'qualifier': 'community fixed decompilation; input hints only', 'lists': lists,
            'common_input_evidence': [{'function': name, 'line': next(i + 1 for i, line in enumerate(lines)
                                    if line.startswith('function ' + name + '('))}
                                    for name in ['GetSwordArtsRequestNew', 'SetSwordArtsPointInfo', 'GetSwordArtsDiffCategory']]}


class Collector:
    def __init__(self, source_dir):
        self.source_dir = source_dir
        self.manifest = []
        for name in ['source-manifest', 'reference-manifest', 'supplemental-manifest']:
            self.manifest.extend(read(source_dir / (name + '.json')))
        verify_manifest(source_dir, self.manifest)
        self.params = {name: {row['Row ID']: row for row in read(source_dir / (name + '.json'))}
                       for name in PARAM_FILES}
        self.source_animations = read(source_dir / 'animations.json')
        self.source_attacks = read(source_dir / 'attacks.json')
        self.armaments = read(source_dir / 'armaments.json')
        wrapper = read(ROOT / 'data/elden-ring-payload.json')
        payload_raw = gzip.decompress(base64.b64decode(wrapper['data']))
        if hashlib.sha256(payload_raw).hexdigest() != wrapper['sha256']:
            raise ValueError('existing weapon payload hash mismatch')
        self.payload = json.loads(payload_raw)
        self.payload_sha256 = wrapper['sha256']
        self.motion = openpyxl.load_workbook(source_dir / 'motion-values.xlsx', read_only=True, data_only=True)
        self.planner = openpyxl.load_workbook(source_dir / 'build-planner.xlsx', read_only=True, data_only=True)
        self.scaling = openpyxl.load_workbook(source_dir / 'scaling-data.xlsx', read_only=True, data_only=True)
        self.compatibility = openpyxl.load_workbook(source_dir / 'ash-compatibility.xlsx', read_only=True, data_only=True)
        self.corrects = sheet_parameters(self.motion, 'AttackElementCorrectParam', 'motion-values.xlsx')
        self.planner_corrects = sheet_parameters(self.planner, 'AttackElementCorrectParam', 'build-planner.xlsx')
        self.graphs = sheet_parameters(self.scaling, 'CalcCorrectGraph', 'scaling-data.xlsx')
        self.sheet_effects = sheet_parameters(self.motion, 'SpEffectParam', 'motion-values.xlsx', header_row=2)
        self.sheet_reinforce = sheet_parameters(self.planner, 'ReinforceParamWeapon', 'build-planner.xlsx')
        self.sheet_attacks = collections.defaultdict(list)
        for index, row in enumerate(self.motion['Ashes of War Attack Data'].values, 1):
            if isinstance(row[2], (int, float)):
                self.sheet_attacks[int(row[2])].append({'row': index, 'weapon': row[0], 'name': row[1],
                                                       'typed_mv': dict(zip(TYPES, row[3:8])),
                                                       'typed_flat': dict(zip(TYPES, row[20:25]))})
        hks_commit = read(source_dir / 'hks-commit.json')
        if hks_commit['sha'] != 'd88d6441f5fccfdd6a5fd10d493309b680181897' or '1.17' not in hks_commit['commit']['message']:
            raise ValueError('HKS commit version evidence changed')
        self.hks = hks_hints((source_dir / 'c0000.hks').read_text(encoding='utf-8-sig'), hks_commit)
        self.behavior_index = collections.defaultdict(list)
        for bid, behavior in self.params['BehaviorParam_PC'].items():
            self.behavior_index[(bid // 300000000, behavior['variationId'], behavior['behaviorJudgeId'])].append(bid)
        self.armament_by_id = {w['id']: w for w in self.armaments.values()}
        self.issues = []
        self.nodes = {}
        self.variant_checks = {}

    def provenance(self, table, key):
        return {'source': table + '.json', 'row_id': key, 'version_scope': '1.17_fixed_parameter_snapshot'}

    def issue(self, reason, **context):
        self.issues.append({'reason': reason, **context})

    def resolve(self, key):
        kind, value = key.split(':', 1)
        index = int(value)
        table = {'behavior': 'BehaviorParam_PC', 'attack': 'AtkParam_Pc', 'bullet': 'Bullet',
                 'effect': 'SpEffectParam', 'final': 'FinalDamageRateParam', 'weapon': 'EquipParamWeapon',
                 'reinforce': 'ReinforceParamWeapon', 'gem': 'EquipParamGem'}.get(kind)
        source_row = self.params.get(table, {}).get(index) if table else None
        fallback = {'effect': self.sheet_effects, 'correct': self.corrects, 'graph': self.graphs}.get(kind, {}).get(index)
        if source_row is None and fallback is None:
            return None
        node = {'kind': kind, 'parameter_id': index, 'edges': [], 'fields': None, 'provenance': [],
                'version_scope': '1.17_fixed_parameter_snapshot' if source_row else fallback['version_scope']}
        if source_row is not None:
            fields = projected(source_row)
            node.update(fields=fields, source_name=source_row['Row Name'])
            node['provenance'].append(self.provenance(table, index))
        else:
            fields = fallback['fields']
            node.update(fields=fields, source_name=fallback['name'])
            node['provenance'].append(fallback['provenance'])
        if source_row is not None and kind == 'effect' and fallback is not None:
            node['reference_sheet_comparison'] = self.compare(fields, fallback['fields'])
            node['reference_sheet_provenance'] = fallback['provenance']
        if kind == 'correct':
            other = self.planner_corrects.get(index)
            node['reference_sheet_comparison'] = self.compare(fields, other['fields']) if other else {'status': 'not_found'}
            node['reference_sheet_provenance'] = other['provenance'] if other else None
        if kind == 'reinforce':
            other = self.sheet_reinforce.get(index)
            node['reference_sheet_comparison'] = self.compare(fields, other['fields']) if other else {'status': 'not_found'}
            node['reference_sheet_provenance'] = other['provenance'] if other else None

        def edge(field, target_kind, allow_zero=False):
            target = fields.get(field)
            if isinstance(target, (int, float)) and not isinstance(target, bool) and target >= (0 if allow_zero else 1):
                node['edges'].append({'field': field, 'node': f'{target_kind}:{int(target)}'})

        if kind == 'behavior':
            target_kind = {0: 'attack', 1: 'bullet', 2: 'effect'}.get(fields.get('refType'))
            if target_kind:
                edge('refId', target_kind)
            else:
                node['unsupported_reference_type'] = fields.get('refType')
        elif kind == 'bullet':
            edge('atkId_Bullet', 'attack')
            for field in ['HitBulletID', 'intervalCreateBulletId']:
                edge(field, 'bullet')
            for field in ['spEffectIDForShooter', *[f'spEffectId{i}' for i in range(5)]]:
                edge(field, 'effect')
            node['num_shoot_is_not_hit_count'] = True
        elif kind == 'attack':
            for field in [f'spEffectId{i}' for i in range(5)]:
                edge(field, 'effect')
            edge('overwriteAttackElementCorrectId', 'correct', allow_zero=True)
            edge('finalDamageRateId', 'final', allow_zero=True)
            node['typed_mv'] = {kind: fields.get('atk' + suffix + 'Correction') for kind, suffix in TYPES.items()}
            node['typed_flat'] = {kind: fields.get('atk' + suffix) for kind, suffix in TYPES.items()}
            nonzero_mv = any(v for v in node['typed_mv'].values() if v is not None)
            nonzero_flat = any(v for v in node['typed_flat'].values() if v is not None)
            node['formula_family_candidate'] = ('weapon_mv_plus_added_base' if nonzero_mv and nonzero_flat else
                                                 'weapon_mv' if nonzero_mv else
                                                 'added_base_with_scaling' if nonzero_flat and fields.get('isAddBaseAtk') else
                                                 'flat_or_special' if nonzero_flat else 'zero_damage_or_effect')
            node['formula_approved'] = False
            node['pvp_final_rate_scope'] = 'motion-values.xlsx/Info!C6,L6'
        elif kind == 'effect':
            for field in EFFECT_LINKS:
                edge(field, 'effect')
            # Behavior IDs need their original category; do not guess a bullet from a number.
            edge('behaviorId', 'behavior')
            node['duration_seconds_raw'] = fields.get('effectEndurance')
            node['stack_group_raw'] = fields.get('spCategory')
            node['priority_raw'] = fields.get('categoryPriority')
            node['activation_and_consumption_verified'] = False
        elif kind == 'gem':
            for field in [f'spEffectId{i}' for i in range(3)] + [f'spEffectId_forAtk{i}' for i in range(3)]:
                edge(field, 'effect')
        elif kind == 'weapon':
            for field in ['residentSpEffectId', 'residentSpEffectId1', 'residentSpEffectId2',
                          'spEffectBehaviorId0', 'spEffectBehaviorId1', 'spEffectBehaviorId2']:
                edge(field, 'effect')
            edge('attackElementCorrectId', 'correct', allow_zero=True)
            for field in fields:
                if field.startswith('correctType_'):
                    edge(field, 'graph', allow_zero=True)
            edge('reinforceTypeId', 'reinforce', allow_zero=True)
        elif kind == 'reinforce':
            node['effect_ids_are_offsets'] = True
        node['edges'].sort(key=lambda e: (e['node'], e['field']))
        return node

    @staticmethod
    def compare(left, right):
        common = sorted(set(left) & set(right))
        mismatches = {key: {'parameter': left[key], 'reference': right[key]} for key in common
                      if left[key] != right[key]}
        return {'status': 'matched_common_fields' if not mismatches else 'mismatch',
                'common_fields': len(common), 'differences': mismatches}

    def normalize_variant(self, ref):
        source = self.source_attacks.get(ref)
        if source is None:
            return None
        refs = source['refs']
        behavior_id = int(ref.split(':')[0])
        behavior = self.params['BehaviorParam_PC'].get(behavior_id)
        roots = []
        if behavior:
            roots.append(f'behavior:{behavior_id}')
        elif not (refs.get('atkParamId') == behavior_id and refs.get('bulletId')):
            self.issue('generated_key_has_no_behavior_row', attack_ref=ref, parsed_id=behavior_id)
        for field, kind in [('atkParamId', 'attack'), ('bulletId', 'bullet')]:
            if isinstance(refs.get(field), int) and refs[field] > 0:
                roots.append(f'{kind}:{refs[field]}')
        roots.extend(f'effect:{i}' for i in refs.get('spEffectIds', []) if i > 0)
        ap = self.params['AtkParam_Pc'].get(refs.get('atkParamId'))
        checks = {'raw_generated_mv': 'no_damage_parameter', 'raw_generated_flat': 'no_damage_parameter',
                  'reference_sheet_mv': 'not_found'}
        if ap:
            typed_mv = {kind: ap['atk' + suffix + 'Correction'] for kind, suffix in TYPES.items()}
            generated_mv = {kind: source.get(kind + 'DamageMV') for kind in TYPES}
            checks['raw_generated_mv'] = 'matched' if typed_mv == generated_mv else 'mismatch'
            if typed_mv != generated_mv:
                self.issue('raw_generated_mv_mismatch', attack_ref=ref, raw=typed_mv, generated=generated_mv)
            typed_flat = {kind: ap['atk' + suffix] for kind, suffix in TYPES.items()}
            generated_flat = {kind: source.get(kind + 'DamageFlat') for kind in TYPES}
            checks['raw_generated_flat'] = 'matched' if typed_flat == generated_flat else 'mismatch'
            if typed_flat != generated_flat:
                self.issue('raw_generated_flat_mismatch', attack_ref=ref, raw=typed_flat, generated=generated_flat)
            sheets = self.sheet_attacks.get(refs['atkParamId'], [])
            checks['reference_sheet_mv'] = ('matched' if any(row['typed_mv'] == typed_mv for row in sheets) else
                                            'mismatch' if sheets else 'not_found')
            if sheets and checks['reference_sheet_mv'] == 'mismatch':
                self.issue('reference_sheet_mv_mismatch', attack_ref=ref, attack_id=refs['atkParamId'])
        else:
            sheets = []
        self.variant_checks[ref] = checks
        return {'source_name': source.get('name'), 'refs': refs, 'weapon_candidates': source.get('weapons', []),
                'source_fields': {k: v for k, v in source.items() if k not in ['refs', 'weapons', 'weaponSkills', 'name']},
                'parameter_roots': sorted(set(roots)), 'validation': checks,
                'provenance': {'source': 'attacks.json', 'key': ref},
                'reference_sheet_rows': [{'source': 'motion-values.xlsx', 'sheet': 'Ashes of War Attack Data',
                                          'row': row['row'], 'weapon_label': row['weapon'], 'action_label': row['name']}
                                         for row in sheets]}

    def compatibility_records(self, skills):
        name_ids = collections.defaultdict(list)
        for skill in skills.values():
            if skill['name_en']:
                name_ids[skill['name_en']].append(skill['skill_id'])
        rows, by_skill = {}, collections.defaultdict(list)
        for index, row in enumerate(self.compatibility['Weapons'].values, 1):
            if not isinstance(row[0], (int, float)):
                continue
            wid = int(row[0])
            if wid in rows:
                raise ValueError(f'duplicate weapon/affinity row {wid}')
            offered = sorted(set(v for v in row[8:] if isinstance(v, str) and v not in ['-', '']))
            aliases = []
            for name in offered:
                ids = name_ids.get(name, [])
                # Use a real ash relationship to disambiguate same-name fixed skills.
                ash_ids = [i for i in ids if skills[str(i)]['availability']['ash_ids']]
                selected = ash_ids if ash_ids else ids
                aliases.append({'source_name': name, 'skill_ids': selected,
                                'resolution': 'unique_ash_skill' if len(selected) == 1 and ash_ids else
                                              'unique_catalog_skill' if len(selected) == 1 else 'ambiguous_or_missing'})
                if len(selected) == 1:
                    by_skill[selected[0]].append(wid)
                elif ids:
                    self.issue('compatibility_same_name_requires_id', row=index, weapon_id=wid,
                               source_name=name, skill_ids=ids)
                else:
                    self.issue('compatibility_name_not_in_catalog', row=index, source_name=name)
            rows[str(wid)] = {'weapon_param_id': wid, 'name_en': row[1], 'category': row[2], 'affinity': row[3],
                              'default_skill_label': row[6], 'can_apply_ash': row[7], 'offered_skills': aliases,
                              'provenance': {'source': 'ash-compatibility.xlsx', 'sheet': 'Weapons', 'row': index},
                              'status': 'reference_sheet_candidate_requires_parameter_crosscheck'}
        return rows, by_skill

    def missing_window_candidates(self, event, users):
        judge = event['params'].get('BehaviorJudgeID')
        if not isinstance(judge, int):
            return []
        variations = {self.armament_by_id[user['id']]['behavior_variation_id'] for user in users
                      if user['type'] == 'Weapon' and user['id'] in self.armament_by_id}
        variations.add(0)
        # Observed ID bank convention narrows candidates; it does not prove execution.
        return sorted({bid for variation in variations for bid in
                       self.behavior_index[(judge // 3000, variation, judge % 3000)]})

    def collect(self):
        skills, animations, actions, variants = {}, {}, {}, {}
        indexed_categories = collections.defaultdict(list)
        for aid, animation in self.source_animations.items():
            for category in animation.get('categories', []):
                indexed_categories[category].append(aid)
        roots = set()
        playable = set()
        for skill_id, param in sorted(self.params['SwordArtsParam'].items()):
            ashes = [a for a in self.payload['ashes'] if a['skill_id'] == skill_id]
            weapons = [w for w in self.payload['weapons'] if w.get('default_skill_id') == skill_id]
            name = param['Row Name']
            classification = ('no_skill' if name == 'No Skill' else 'player_linked_candidate' if ashes or weapons else
                              'internal_or_unclassified')
            if classification == 'player_linked_candidate':
                playable.add(skill_id)
            skill_roots = {f'gem:{int(a["id"].split(":")[1])}' for a in ashes}
            skill_roots.update(f'weapon:{w["game_id"]}' for w in weapons if w['game_id'] in self.params['EquipParamWeapon'])
            section = f"a{600 + param['swordArtsTypeNew']:03d}"
            category_candidates = sorted(indexed_categories[name]) if name and name != 'No Skill' else []
            # Keep shared names and cross-section reuse as candidates, never as extra hits.
            skill = {'skill_id': skill_id, 'name_en': name,
                     'name_ja': self.payload['skills'].get(str(skill_id), {}).get('name_ja'),
                     'classification': classification, 'game_version': '1.17',
                     'parameter_fields': projected(param), 'hks_arts_type': param['swordArtsTypeNew'],
                     'nominal_animation_section': section, 'provenance': self.provenance('SwordArtsParam', skill_id),
                     'availability': {'ash_ids': [a['id'] for a in ashes], 'default_weapon_ids': [w['id'] for w in weapons],
                                      'default_weapon_param_ids': [w['game_id'] for w in weapons],
                                      'compatibility_variant_ids': []},
                     'costs': {button: {'fp_raw': param['useMagicPoint_' + button],
                                       'fp': param['useMagicPoint_' + button] if param['useMagicPoint_' + button] >= 0 else None,
                                       'sentinel': param['useMagicPoint_' + button] < 0}
                               for button in ['L1', 'L2', 'R1', 'R2']},
                     'cost_scope': 'parameter_base_cost_before_equipment_reductions; input binding still requires validation',
                     'hks_candidates': {key: param['swordArtsTypeNew'] in value['members']
                                        for key, value in self.hks['lists'].items()},
                     'animation_ids': category_candidates, 'actions': [],
                     'observed_parameter_roots': [], 'candidate_effect_roots': [], 'candidate_behavior_roots': [],
                     'validation': {'collection': 'collected', 'input_mapping': 'candidate',
                                    'formula': 'unverified', 'damage': 'unmeasured', 'timing': 'unmeasured',
                                    'simulator_ready': False, 'blockers': []}}
            if classification != 'no_skill' and not category_candidates:
                self.issue('no_named_animation_mapping', skill_id=skill_id, section=section)
            # Buff names are discovery evidence only. Activation is not proved by an effect name.
            for eid, effect in self.sheet_effects.items():
                if name and re.match(r'^\[AoW\] ' + re.escape(name) + r'(?:$|\s*[-(])', effect['name'] or ''):
                    skill['candidate_effect_roots'].append(f'effect:{eid}')
            for aid in category_candidates:
                source = self.source_animations[aid]
                if source.get('version') != '1.17':
                    raise ValueError(f'animation baseline changed: {aid}')
                if aid not in animations:
                    animations[aid] = {'name': source['name'], 'section': source['section'],
                                       'source_key': source.get('sourceKey'), 'labels': source.get('labels', []),
                                       'users': source['users'], 'categories': source['categories'],
                                       'source_events': source['events'], 'cancel_from_types': source.get('cancelFrom', []),
                                       'provenance': {'source': 'animations.json', 'key': aid},
                                       'time_scope': '30_source_frames_per_second_display; real_transition_timing_unverified'}
                action_id = f'{skill_id}:{aid}'
                skill['actions'].append(action_id)
                action = {'action_id': action_id, 'skill_id': skill_id, 'animation_id': aid,
                          'source_label': source['name'], 'input_sequence': None, 'fp_total': None,
                          'hit_count': None, 'cycle_seconds': None, 'damage': None, 'damage_per_fp': None, 'dps': None,
                          'fp_branch_label': 'insufficient_fp' if re.search(r'No FP|without FP', source['name'], re.I) else
                                             'sufficient_fp' if 'with FP' in source['name'] else 'unspecified',
                          'windows': [], 'parameter_roots': [], 'candidate_parameter_roots': [],
                          'observed_effect_windows': source['events'].get('spEffects', []),
                          'readiness': 'source_candidate_only', 'blockers': ['input_selector_graph_unverified',
                                        'in_game_validation_missing', 'time_cycle_unverified']}
                action_roots = {f'effect:{event["id"]}' for event in action['observed_effect_windows'] if event['id'] > 0}
                for index, event in enumerate(source['events'].get('activeFrames', [])):
                    refs = event['params'].get('AtkParamID', [])
                    window = {'index': index, 'source_type': event['type'], 'source_range': event['range'],
                              'source_params': event['params'], 'candidate_refs': refs, 'weapon_bindings': {},
                              'component_mode': 'alternatives_and_components_require_selection; never_sum_all_candidates'}
                    if not refs:
                        action['blockers'].append('window_without_generated_attack_refs')
                        candidates = self.missing_window_candidates(event, source['users'])
                        window['candidate_behavior_nodes'] = [f'behavior:{bid}' for bid in candidates]
                        window['candidate_basis'] = 'observed_behavior_ID_bank_and_weapon_variation; execution_unverified'
                        action['candidate_parameter_roots'].extend(window['candidate_behavior_nodes'])
                        self.issue('window_without_generated_attack_refs', skill_id=skill_id, action_id=action_id,
                                   window=index, has_weapon_users=bool(source['users']),
                                   candidate_behavior_nodes=window['candidate_behavior_nodes'])
                    for ref in refs:
                        if ref not in variants:
                            normalized = self.normalize_variant(ref)
                            if normalized is None:
                                self.issue('missing_generated_attack_ref', skill_id=skill_id, action_id=action_id, ref=ref)
                                continue
                            variants[ref] = normalized
                        action_roots.update(variants[ref]['parameter_roots'])
                    user_names = sorted({user['name'] for user in source['users'] if user['type'] == 'Weapon'})
                    for weapon in user_names:
                        window['weapon_bindings'][weapon] = select_bindings(refs, self.source_attacks, weapon)
                    action['windows'].append(window)
                if not source['users']:
                    action['blockers'].append('source_has_no_weapon_users')
                if source['events'].get('gradients'):
                    action['blockers'].append('speed_gradient_requires_integration')
                action['parameter_roots'] = sorted(action_roots)
                action['candidate_parameter_roots'] = sorted(set(action['candidate_parameter_roots']))
                action['blockers'] = sorted(set(action['blockers']))
                actions[action_id] = action
                skill_roots.update(action_roots)
                skill['candidate_behavior_roots'].extend(action['candidate_parameter_roots'])
            skill['observed_parameter_roots'] = sorted(skill_roots)
            roots.update(skill_roots)
            roots.update(skill['candidate_effect_roots'])
            skill['candidate_behavior_roots'] = sorted(set(skill['candidate_behavior_roots']))
            roots.update(skill['candidate_behavior_roots'])
            skills[str(skill_id)] = skill
        compatibility, compat_index = self.compatibility_records(skills)
        # Bring in every actual weapon affinity's scaling, reinforcement, and status parameters.
        for wid in compatibility:
            if int(wid) in self.params['EquipParamWeapon']:
                roots.add(f'weapon:{wid}')
        weapon_variants = {}
        for weapon in self.armaments.values():
            for affinity, variant in weapon['affinity'].items():
                wid = variant['id']
                roots.add(f'weapon:{wid}')
                weapon_variants[str(wid)] = {
                    'weapon_param_id': wid, 'base_weapon_id': weapon['id'], 'name_en': weapon['name'],
                    'category': weapon['category'], 'affinity': affinity,
                    'default_skill_id': weapon['default_skill_id'], 'allow_ash_of_war': weapon['allow_ash_of_war'],
                    'requirements': weapon['requirements'], 'behavior_variation_id': weapon['behavior_variation_id'],
                    'variant_fields': variant, 'max_upgrade': len(weapon['upgrade_costs']),
                    'provenance': {'source': 'armaments.json', 'key': weapon['name'], 'affinity': affinity}}
        ammo_parameters = {}
        for index, row in enumerate(self.motion['AmmoData'].values, 1):
            if not isinstance(row[2], (int, float)):
                continue
            wid = int(row[2])
            roots.add(f'weapon:{wid}')
            ammo_parameters[str(wid)] = {'weapon_param_id': wid, 'source_name': row[1], 'type': row[0],
                                        'provenance': {'source': 'motion-values.xlsx', 'sheet': 'AmmoData', 'row': index}}
        for index, row in enumerate(self.motion['Ammo Attack Data'].values, 1):
            if isinstance(row[2], (int, float)):
                roots.add(f'attack:{int(row[2])}')
        for skill in skills.values():
            skill['availability']['compatibility_variant_ids'] = sorted(set(compat_index[skill['skill_id']]))
        self.nodes, missing = walk_graph(roots, self.resolve)
        # All upgrade levels for referenced weapon reinforcement families, including baseAtkRate.
        reinforcement_ids = {int(k.split(':')[1]) for k in self.nodes if k.startswith('reinforce:')}
        for family in sorted(reinforcement_ids):
            for upgrade in range(26):
                key = f'reinforce:{family + upgrade}'
                if family + upgrade in self.params['ReinforceParamWeapon']:
                    self.nodes[key] = self.resolve(key)
        for action in actions.values():
            deps = reachable(action['parameter_roots'] + action['candidate_parameter_roots'], self.nodes)
            action['dependency_nodes'] = deps
            action['missing_dependency_nodes'] = [key for key in deps if key not in self.nodes]
            if action['missing_dependency_nodes']:
                action['blockers'].append('missing_parameter_dependency')
            action['reference_only_dependency_nodes'] = [key for key in deps if key in self.nodes and
                                                        self.nodes[key]['version_scope'] != '1.17_fixed_parameter_snapshot']
            if action['reference_only_dependency_nodes']:
                action['blockers'].append('reference_sheet_version_bridge_unverified')
            action['blockers'] = sorted(set(action['blockers']))
        topology_cache = {}
        for action in actions.values():
            for window in action['windows']:
                for binding in window['weapon_bindings'].values():
                    bullet_roots = tuple(sorted({f"bullet:{variants[ref]['refs']['bulletId']}"
                                                 for ref in binding['matched_refs']
                                                 if variants[ref]['refs'].get('bulletId')}))
                    if bullet_roots:
                        if bullet_roots not in topology_cache:
                            topology_cache[bullet_roots] = bullet_topology(bullet_roots, self.nodes)
                        binding['bullet_topology'] = topology_cache[bullet_roots]
        for skill in skills.values():
            deps = reachable(skill['observed_parameter_roots'] + skill['candidate_effect_roots'] +
                             skill['candidate_behavior_roots'], self.nodes)
            skill['missing_dependency_nodes'] = [key for key in deps if key not in self.nodes]
            kinds = {self.nodes[key]['kind'] for key in deps if key in self.nodes}
            skill['components_observed'] = sorted(kinds & {'attack', 'bullet', 'effect'})
            skill['mechanic_class_candidate'] = ('mixed_weapon_projectile' if 'attack' in kinds and 'bullet' in kinds else
                                                 'weapon_attack' if 'attack' in kinds else
                                                 'effect_or_movement_requires_semantic_classification')
            skill['non_attack_damage'] = None
            skill['validation']['blockers'] = sorted({blocker for aid in skill['actions'] for blocker in actions[aid]['blockers']}
                                                     | ({'no_named_animation_mapping'} if not skill['actions'] and
                                                        skill['classification'] != 'no_skill' else set())
                                                     | ({'missing_parameter_dependency'} if skill['missing_dependency_nodes'] else set()))
            if skill['candidate_effect_roots']:
                skill['validation']['blockers'].append('named_effect_activation_not_verified')
            if skill['classification'] == 'no_skill':
                skill['validation']['input_mapping'] = 'not_applicable'
        self.check_behavior_routes(variants)
        # Every mismatch reaches the action that uses it; unrelated actions keep their own status.
        for action in actions.values():
            refs = {ref for window in action['windows'] for ref in window['candidate_refs'] if ref in variants}
            if any(variants[ref]['behavior_connection'] == 'unresolved_generated_connection' for ref in refs):
                action['blockers'].append('unresolved_generated_behavior_connection')
            if any('mismatch' in variants[ref]['validation'].values() for ref in refs):
                action['blockers'].append('attack_coefficient_mismatch')
            action['blockers'] = sorted(set(action['blockers']))
        for skill in skills.values():
            skill['validation']['blockers'] = sorted(set(skill['validation']['blockers']) |
                                                     {b for aid in skill['actions'] for b in actions[aid]['blockers']})
        self.check_compatibility_params(compatibility)
        cases = self.measurement_cases(skills, actions, compatibility)
        missing_references = {key: {
            'referring_parameters': [{'node': parent, 'field': edge['field']}
                                      for parent, row in self.nodes.items() for edge in row['edges'] if edge['node'] == key],
            'dependent_actions': sorted(aid for aid, action in actions.items() if key in action['missing_dependency_nodes']),
            'dependent_skill_ids': sorted(skill['skill_id'] for skill in skills.values() if key in skill['missing_dependency_nodes']),
            'resolution': 'no_verified_row_in_collected_sources'} for key in missing}
        audit = {'schema_version': 2, 'baseline': '1.17', 'scope': 'all_rows_in_fixed_SwordArtsParam',
                 'counts': {'catalog_rows': len(skills), 'player_linked_candidates': len(playable),
                            'player_candidates_with_animations': sum(bool(skills[str(i)]['animation_ids']) for i in playable),
                            'classifications': dict(collections.Counter(s['classification'] for s in skills.values())),
                            'skill_animation_actions': len(actions), 'unique_animations': len(animations),
                            'generated_attack_variants': len(variants), 'parameter_nodes': len(self.nodes),
                            'source_files': len(self.manifest), 'weapon_variants': len(weapon_variants),
                            'ammo_parameters': len(ammo_parameters),
                            'nodes_by_kind': dict(collections.Counter(n['kind'] for n in self.nodes.values())),
                            'missing_parameter_nodes': len(missing), 'issues': len(self.issues),
                            'issue_reasons': dict(collections.Counter(i['reason'] for i in self.issues)),
                            'unmapped_windows_with_weapon_users': sum(i['reason'] == 'window_without_generated_attack_refs' and
                                                                      i['has_weapon_users'] for i in self.issues),
                            'unmapped_windows_without_weapon_users': sum(i['reason'] == 'window_without_generated_attack_refs' and
                                                                         not i['has_weapon_users'] for i in self.issues),
                            'player_candidates_with_missing_dependencies': sum(bool(s['missing_dependency_nodes']) for s in skills.values()
                                                                              if s['classification'] == 'player_linked_candidate'),
                            'planned_measurement_cases': len(cases), 'measured_cases': 0,
                            'simulator_ready_skills': 0},
                 'missing_parameter_nodes': missing, 'cycle_witnesses': cycle_witnesses(self.nodes),
                 'variant_validation_counts': {check: dict(collections.Counter(v[check] for v in self.variant_checks.values()))
                                               for check in ['raw_generated_mv', 'raw_generated_flat', 'reference_sheet_mv']},
                 'source_hash_drift': [s['file'] for s in self.manifest if s.get('sha256') and
                                       s['sha256'] != s['actual_sha256']],
                 'source_limitations': ['HKS_commit_declares_1.17_but_header_stale_community_edits',
                                        'full_TAE_sample_endpoints_return_404_alternative_repository_is_1.14',
                                        '1.17_to_1.17.1_parameter_diff_unverified',
                                        'reference_workbooks_can_share_the_same_game_data_origin'],
                 'ready_for_source_driven_implementation_review': True, 'fully_verified_simulation': False}
        metadata = {'schema_version': 2, 'game_version': '1.17',
                    'existing_weapon_payload_sha256': self.payload_sha256,
                    'units': {'mv': 'percent; 100=1', 'animation': 'source frames; display uses 30/s',
                              'effect_endurance': 'seconds_raw; sentinel interpretation required'},
                    'contracts': ['null is unknown, not zero', 'no_skill and non_damage are distinct',
                                  'alternative references are never automatically summed',
                                  'numShoot is spawn count, not hits against a target',
                                  'named effect matches are candidate discovery, not activation proof',
                                  'PvP final rates do not apply to PvE',
                                  'reference-only parameters retain separate version scope',
                                  'parameter collection does not prove damage or timing'],
                    'projection_excludes': list(VISUAL_PREFIXES), 'hks': self.hks,
                    'formula_guidance': {'source': 'skill-damage-explanation.txt',
                                         'scope': 'historical reference mentions 1.04.1; numeric examples not adopted'}}
        return {'metadata': metadata, 'skills': skills, 'actions': actions, 'animations': animations,
                'attack-variants': variants, 'parameters': self.nodes, 'compatibility': compatibility,
                'measurement-cases': cases, 'issues': self.issues, 'audit': audit,
                'source-manifest': self.manifest, 'weapon-variants': weapon_variants, 'ammo-parameters': ammo_parameters,
                'missing-references': missing_references}

    def check_behavior_routes(self, variants):
        for ref, variant in variants.items():
            behavior_id = int(ref.split(':')[0])
            behavior = self.params['BehaviorParam_PC'].get(behavior_id)
            if not behavior:
                variant['behavior_connection'] = 'generated_direct_attack_alias'
                continue
            target_kind = {0: 'attack', 1: 'bullet', 2: 'effect'}.get(behavior['refType'])
            root = f"{target_kind}:{behavior['refId']}"
            deps = set(reachable([root], self.nodes)) if behavior['refId'] > 0 else set()
            expected = {f"{kind}:{variant['refs'][field]}" for field, kind in [('atkParamId', 'attack'), ('bulletId', 'bullet')]
                        if variant['refs'].get(field)}
            connection = ('source_placeholder_without_damage' if not expected else
                          'root_or_descendant' if expected <= deps else 'unresolved_generated_connection')
            variant['behavior_connection'] = connection
            if connection == 'unresolved_generated_connection':
                self.issue('unresolved_generated_behavior_connection', attack_ref=ref, root=root,
                           unreachable=sorted(expected - deps))

    def check_compatibility_params(self, records):
        for key, row in records.items():
            param = self.params['EquipParamWeapon'].get(int(key))
            if param is None:
                row['parameter_validation'] = 'weapon_row_missing'
                self.issue('compatibility_weapon_parameter_missing', weapon_param_id=int(key))
            else:
                row['parameter_validation'] = 'parameter_found; per-ash exceptions still require review'
                row['default_skill_id_raw'] = param['swordArtsParamId']

    @staticmethod
    def measurement_cases(skills, actions, compatibility):
        cases = []
        for skill in skills.values():
            if skill['classification'] != 'player_linked_candidate':
                continue
            compatible = skill['availability']['compatibility_variant_ids']
            defaults = skill['availability']['default_weapon_param_ids']
            legal_ids = set(compatible + defaults)
            for aid in skill['actions']:
                action = actions[aid]
                users = sorted({weapon for window in action['windows'] for weapon, binding in window['weapon_bindings'].items()
                                if binding['resolution'] == 'unique_or_exact_alias'})
                candidates = [int(key) for key in compatibility if int(key) in legal_ids and
                              (not users or compatibility[key]['name_en'] in users)]
                # A sample per observed animation; the full compatibility matrix is retained separately.
                wid = min(candidates) if candidates else None
                cases.append({'case_id': aid, 'skill_id': skill['skill_id'], 'animation_id': action['animation_id'],
                              'weapon_param_id_candidate': wid,
                              'availability_evidence': compatibility[str(wid)]['provenance'] if wid is not None else None,
                              'game_version': None, 'input_sequence': None, 'upgrade': None, 'attributes': None,
                              'hands': None, 'target': None, 'journey': None, 'equipment_effects': None,
                              'distance': None, 'fp_before': None, 'fp_after': None, 'observed_damage': None,
                              'trials': [], 'status': 'awaiting_controlled_measurement'})
        return cases


def write_outputs(out, data):
    out.mkdir(parents=True, exist_ok=True)
    for name, value in data.items():
        (out / (name + '.json')).write_bytes(encode(value) + b'\n')
    with (out / 'coverage.csv').open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(['skill_id', 'name_ja', 'name_en', 'classification', 'fp_L2', 'fp_R1', 'fp_R2',
                         'animations', 'missing_dependencies', 'simulator_ready', 'blockers'])
        for skill in sorted(data['skills'].values(), key=lambda row: row['skill_id']):
            writer.writerow([skill['skill_id'], skill['name_ja'], skill['name_en'], skill['classification'],
                             *[skill['costs'][button]['fp'] for button in ['L2', 'R1', 'R2']],
                             len(skill['animation_ids']), len(skill['missing_dependency_nodes']),
                             skill['validation']['simulator_ready'], '|'.join(skill['validation']['blockers'])])
    runtime = {name.replace('-', '_'): data[name] for name in data if name not in ['measurement-cases', 'issues', 'audit', 'missing-references']}
    raw = encode(runtime)
    payload = {'encoding': 'gzip-base64', 'sha256': hashlib.sha256(raw).hexdigest(), 'raw_bytes': len(raw),
               'data': base64.b64encode(gzip.compress(raw, compresslevel=9, mtime=0)).decode()}
    (out / 'payload.json').write_bytes(encode(payload) + b'\n')
    index = {'schema_version': 2, 'files': {name: {'sha256': hashlib.sha256((out / name).read_bytes()).hexdigest(),
                                                 'bytes': (out / name).stat().st_size}
                                           for name in sorted([name + '.json' for name in data] + ['payload.json', 'coverage.csv'])}}
    (out / 'index.json').write_bytes(encode(index) + b'\n')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--sources', type=Path, required=True)
    parser.add_argument('--out', type=Path, default=ROOT / 'data/elden-ring-skills/all')
    args = parser.parse_args()
    data = Collector(args.sources).collect()
    write_outputs(args.out, data)
    print(json.dumps(data['audit'], ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
